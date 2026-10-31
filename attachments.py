"""Extract readable text from Word (.docx) and Excel (.xlsx) files.

The full extractable text is included or the upload is rejected: never silently
truncate documents to make them fit in an LLM context window.
"""
from __future__ import annotations

from io import BytesIO
from pathlib import PurePath

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

MAX_FILES = 5
MAX_FILE_BYTES = 8_000_000
MAX_EXTRACTED_CHARS = 13_000


def extract_word(content: bytes) -> str:
    """Read paragraphs, tables, headers and footers from a DOCX file."""
    document = Document(BytesIO(content))
    lines: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            if block.text.strip():
                lines.append(block.text.strip())
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
    for section in document.sections:
        for item in (section.header, section.footer):
            for paragraph in item.paragraphs:
                if paragraph.text.strip():
                    lines.append(paragraph.text.strip())
            for table in item.tables:
                for row in table.rows:
                    cells = [cell.text.strip().replace("\n", " / ") for cell in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
    if not lines:
        raise ValueError("O Word não contém texto extraível de parágrafos ou tabelas.")
    return "\n".join(lines)


def extract_excel(content: bytes) -> str:
    """Read ALL sheets and non-empty cells, including formulas when present.

    Cached formula values are included when the Excel file contains them.
    Excel charts, images and formatting are not interpreted.
    """
    formulas = load_workbook(BytesIO(content), read_only=True, data_only=False)
    cached = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        lines: list[str] = []
        for sheet in formulas:
            values = cached[sheet.title]
            lines.append(f"[FOLHA: {sheet.title}]")
            # zip allows us to read formula and cached-value views of all rows
            # without materializing the entire workbook in memory.
            for row_formula, row_value in zip(sheet.iter_rows(), values.iter_rows()):
                items = []
                for cell, value in zip(row_formula, row_value):
                    if cell.value is None:
                        continue
                    position = f"{get_column_letter(cell.column)}{cell.row}"
                    content_value = str(cell.value).replace("\n", " ").strip()
                    if cell.data_type == "f" and value.value is not None:
                        content_value += f" [valor guardado: {value.value}]"
                    items.append(f"{position}={content_value}")
                if items:
                    lines.append(" | ".join(items))
        if not any(line for line in lines if not line.startswith("[FOLHA:")):
            raise ValueError("O Excel não tem células preenchidas.")
        return "\n".join(lines)
    finally:
        formulas.close()
        cached.close()


def build_evidence(files: list[tuple[str, bytes]]) -> str:
    """Create one evidence string for the debate, keeping EVERY extracted cell/line."""
    if len(files) > MAX_FILES:
        raise ValueError(f"Podes anexar no máximo {MAX_FILES} ficheiros por debate.")
    extracted: list[str] = []
    total = 0
    for raw_name, content in files:
        name = PurePath(raw_name).name.replace("\n", " ").replace("\r", " ")
        if len(content) > MAX_FILE_BYTES:
            raise ValueError(f"{name}: ultrapassa o limite de 8 MB por ficheiro.")
        suffix = name.lower().rsplit(".", 1)[-1]
        try:
            if suffix == "docx":
                text = extract_word(content)
            elif suffix == "xlsx":
                text = extract_excel(content)
            else:
                raise ValueError("Só são aceites ficheiros .docx e .xlsx.")
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"Não foi possível ler {name}: ficheiro inválido ou danificado.") from exc
        part = f"\n### ANEXO: {name}\n{text}\n"
        total += len(part)
        if total > MAX_EXTRACTED_CHARS:
            raise ValueError(
                "Os documentos contêm demasiado texto para este modelo local "
                f"({total:,} caracteres extraídos; limite {MAX_EXTRACTED_CHARS:,}). "
                "Para manter exatamente 7 chamadas e não omitir dados, usa ficheiros mais pequenos."
            )
        extracted.append(part)
    return "\n".join(extracted)
